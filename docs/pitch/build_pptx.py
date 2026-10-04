"""Build a static PPTX with one image per slide, manual advancement and French notes.
Run via uv --no-project --with python-pptx python build_pptx.py in docs/pitch.
"""
import argparse
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu
HERE=Path(__file__).resolve().parent
WIDTH,HEIGHT=Emu(12192000),Emu(6858000)
slides=json.loads((HERE/'slides/slides.json').read_text())
parser=argparse.ArgumentParser()
parser.add_argument('--refresh-images',action='store_true',help='Update slide PNGs in an existing PPTX, preserving all other contents, including edited notes.')
args=parser.parse_args()
output=HERE/'Club-des-Affaires-pitch.pptx'
temporary=HERE/'.build/pitch-export.pptx'
temporary.parent.mkdir(exist_ok=True)
if args.refresh_images:
    original=output.read_bytes()
    deck=Presentation(BytesIO(original))
    if len(deck.slides)!=len(slides):
        raise SystemExit('Slide count changed; refusing to overwrite an edited presentation.')
    replacements={}
    for info,slide in zip(slides,deck.slides):
        if len(slide.shapes)!=1 or not hasattr(slide.shapes[0],'image'):
            raise SystemExit('Slide structure changed; refusing to overwrite an edited slide.')
        shape=slide.shapes[0]
        image_part=slide.part.related_part(shape._element.blipFill.blip.rEmbed)
        name=str(image_part.partname).lstrip('/')
        data=(HERE/info['image']).read_bytes()
        if name in replacements and replacements[name]!=data:
            raise SystemExit('Shared image has conflicting replacements.')
        replacements[name]=data
    with ZipFile(BytesIO(original)) as source:
        changed={name:data for name,data in replacements.items() if source.read(name)!=data}
        with ZipFile(temporary,'w') as target:
            target.comment=source.comment
            for entry in source.infolist():
                target.writestr(entry,changed.get(entry.filename,source.read(entry.filename)))
    if output.read_bytes()!=original:
        temporary.unlink()
        raise SystemExit('PowerPoint changed during export; original file preserved. Close/save it and retry.')
    (HERE/'.build/pitch-before-image-refresh.pptx').write_bytes(original)
    temporary.replace(output)
    print(f'{output.name}: {len(changed)} updated images; all notes, XML and other contents preserved')
    raise SystemExit(0)
deck=Presentation()
deck.slide_width,deck.slide_height=WIDTH,HEIGHT
for info in slides:
    slide=deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(HERE/info['image']),0,0,WIDTH,HEIGHT)
    transition=OxmlElement('p:transition')
    transition.set('advClick','1')
    slide._element.insert_element_before(transition,'p:timing','p:extLst')
    slide.notes_slide.notes_text_frame.text=info['notes']
# Atomic replacement also avoids leaving a partial file if export fails.
deck.save(temporary)
temporary.replace(output)
print(f'{output.name}: {len(slides)} static slides and French notes')
