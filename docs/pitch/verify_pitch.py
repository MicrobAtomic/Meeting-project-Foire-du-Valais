"""Validate static exports, the requested screen order and notes without changing the app."""
import argparse
import json
from pathlib import Path
import posixpath
import re
from zipfile import ZipFile
from lxml import etree
from PIL import Image
from pptx import Presentation
HERE=Path(__file__).resolve().parent
parser=argparse.ArgumentParser()
parser.add_argument('--preserved-pptx',type=Path,help='Original edited PPTX: check that everything except slide PNGs was preserved.')
parser.add_argument('--inserted-slide',type=int,help='Position of the newly inserted slide when comparing with an original PPTX.')
parser.add_argument('--updated-notes',type=int,nargs='*',default=[])
args=parser.parse_args()
slides=json.loads((HERE/'slides/slides.json').read_text())
story=json.loads((HERE/'story.json').read_text())
captures=json.loads((HERE/'screens/manifest.json').read_text())
def slide_part_names(archive):
    # python-pptx renumbers slide part names in memory when reading their order.
    # Compare the actual package paths to verify that existing parts were retained.
    ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
    root=etree.fromstring(archive.read('ppt/presentation.xml'))
    relationships=etree.fromstring(archive.read('ppt/_rels/presentation.xml.rels'))
    targets={row.get('Id'):posixpath.normpath('ppt/'+row.get('Target')) if not row.get('Target').startswith('/') else row.get('Target').lstrip('/') for row in relationships}
    return [targets[row.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')] for row in root.findall('p:sldIdLst/p:sldId',ns)]
deck=Presentation(HERE/'Club-des-Affaires-pitch.pptx')
assert len(slides)==len(deck.slides)==27
assert len(story)==19 and sum(s['seconds'] for s in story)==118
assert len(re.findall(rb'/Type\s*/Page\b',(HERE/'Club-des-Affaires-pitch.pdf').read_bytes()))==27
assert [(row['id'],row['screen']) for row in story[14:18]]==[
    ('admin-dashboard','screens/admin/000.png'),('admin-prepare','screens/admin/001.png'),
    ('admin-tables','screens/admin/002.png'),('admin-badges','screens/admin/003.png')]
assert story[0]['screen']=='screens/web/000.png'
assert [(r['id'],r['screen']) for r in story[2:6]]==[
    ('web-login','screens/web/001.png'),('web-dashboard','screens/web/002.png'),
    ('web-album','screens/web/003.png'),('web-intros','screens/web/004.png')]
used=[r['screen'] for r in story if r['screen']]
expected=[f'screens/{name}/{frame["file"]}' for name,frames in captures['scenes'].items() for frame in frames]
assert set(used)==set(expected) and len(used)==len(set(used))==17
assert captures['errors']==[]
clock=0
for row in story:
    assert row['start']==clock and row['end']==clock+row['seconds'];clock=row['end']
for i,(row,slide) in enumerate(zip(slides,deck.slides)):
    assert Image.open(HERE/row['image']).size==(1920,1080)
    if not args.preserved_pptx:
        assert slide.notes_slide.notes_text_frame.text==row['notes']
    assert len(slide.shapes)==1 and slide.shapes[0].image.blob==(HERE/row['image']).read_bytes()
    if row['id']:
        script=story[i]
        assert row['id']==script['id'] and row['notes'].endswith(script['text'])
        assert row['seconds']==script['seconds'] and row['chapter']==script['chapter']
        assert script['text'] in (HERE.parent/'PITCH.md').read_text()
        assert [s['source'] for s in row['screens']]==([script['screen']] if script['screen'] else [])
    transition=slide._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}transition')
    assert transition is not None and transition.get('advClick','1') in ('1','true') and transition.get('advTm') is None
    assert slide._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}timing') is None
with ZipFile(HERE/'Club-des-Affaires-pitch.pptx') as archive:
    assert not any(name.endswith(('.gif','.mp4')) for name in archive.namelist())
    assert all(name.endswith('.png') for name in archive.namelist() if name.startswith('ppt/media/'))
    if args.preserved_pptx:
        with ZipFile(args.preserved_pptx) as original:
            if not args.inserted_slide:
                assert archive.namelist()==original.namelist()
            media={str(slide.part.related_part(slide.shapes[0]._element.blipFill.blip.rEmbed).partname).lstrip('/') for slide in deck.slides}
            allowed=set(media)
            if args.inserted_slide:
                previous=Presentation(args.preserved_pptx)
                assert len(previous.slides)+1==len(deck.slides)
                inserted=args.inserted_slide-1
                new=deck.slides[inserted]
                current_parts=slide_part_names(archive)
                previous_parts=slide_part_names(original)
                added={current_parts[inserted],str(new.notes_slide.part.partname).lstrip('/')}
                added|={name.replace('/slides/','/slides/_rels/').replace('/notesSlides/','/notesSlides/_rels/')+'.rels' for name in list(added)}
                added|=media-set(original.namelist())
                assert set(archive.namelist())-set(original.namelist())==added
                allowed|={'[Content_Types].xml','ppt/presentation.xml','ppt/_rels/presentation.xml.rels','docProps/app.xml'}
                for index,old in enumerate(previous.slides):
                    current=deck.slides[index+(index>=inserted)]
                    assert current_parts[index+(index>=inserted)]==previous_parts[index]
                    if index+1 in args.updated_notes:
                        allowed.add(str(old.notes_slide.part.partname).lstrip('/'))
                        assert current.notes_slide.notes_text_frame.text==slides[index+(index>=inserted)]['notes']
                    else:
                        assert current.notes_slide.notes_text_frame.text==old.notes_slide.notes_text_frame.text
                assert new.notes_slide.notes_text_frame.text==slides[inserted]['notes']
            changed=[name for name in original.namelist() if archive.read(name)!=original.read(name)]
            assert set(changed)<=allowed, f'Unexpected changes to the edited PPTX: {changed}'
            print(f'PASS: insertion/image updates verified; all other PPTX contents preserved byte for byte; updated notes {args.updated_notes}.')
print('PASS: 19 manual slides + divider + 7 appendices; 17 unique fixed screens; cover and web order; notes; 27 PDF pages; 118s plan.')
